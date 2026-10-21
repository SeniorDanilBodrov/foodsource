"""Generate a readable XLSX supplier selection without a Node.js runtime."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


HEADERS = [
    "Поставщик", "Тип", "Категории", "Товары", "Город", "Регионы",
    "Мин. заказ", "Цена", "Доставка", "Документы", "Телефон", "Email",
    "Адрес", "Контакт", "Сайт", "Источник", "Обновлено", "Полнота, %",
]


def export_suppliers_xlsx(suppliers: list[dict], output_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Поставщики"
    sheet.sheet_view.showGridLines = False

    sheet.merge_cells("A1:R1")
    sheet["A1"] = "Выборка поставщиков FoodSource"
    sheet["A1"].font = Font(name="Arial", size=15, bold=True, color="171B19")
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 28

    sheet.append(HEADERS)
    for supplier in suppliers:
        sheet.append([
            supplier.get("name", ""),
            supplier.get("type", ""),
            "; ".join(supplier.get("categories") or []),
            "; ".join(supplier.get("products") or []),
            supplier.get("city", ""),
            "; ".join(supplier.get("regions") or []),
            supplier.get("minimumOrder", ""),
            supplier.get("priceRange", ""),
            supplier.get("delivery", ""),
            "; ".join(supplier.get("certificates") or []) or "Не указаны",
            supplier.get("phone", ""),
            supplier.get("email", ""),
            supplier.get("address", ""),
            supplier.get("contact", "") or "Не указан",
            supplier.get("website", ""),
            supplier.get("source", ""),
            supplier.get("updated", ""),
            supplier.get("completeness", 0),
        ])

    header_fill = PatternFill("solid", fgColor="171B19")
    for cell in sheet[2]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=3, max_row=sheet.max_row):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="222622")
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    widths = [26, 22, 25, 34, 23, 32, 18, 22, 32, 26, 20, 26, 35, 27, 30, 30, 14, 12]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:R{sheet.max_row}"
    table = Table(displayName="SuppliersExport", ref=f"A2:R{sheet.max_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False,
        showLastColumn=False, showRowStripes=True, showColumnStripes=False,
    )
    sheet.add_table(table)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
